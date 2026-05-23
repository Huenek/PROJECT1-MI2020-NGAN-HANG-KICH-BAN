import os
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment

script_dir = '/Users/doanvinhnhan/doan1/scriptwriting/raw_script'
topic_file = '/Users/doanvinhnhan/doan1/Scriptwriting/topic.txt'
output_file = '/Users/doanvinhnhan/doan1/scriptwriting/aggregated_scripts.xlsx'

filename_to_topic_idx = {
    'bayes_spam_filter': 1,
    'ber_transmission': 2,
    'system_reliability': 3,
    'decision_tree_investment': 4,
    'exponential_server_response': 5,
    'normal_dist_machining': 6,
    'gaussian_noise_signal': 7,
    'covariance_portfolio': 8,
    '3d_emission_diffusion': 9,
    'robot_positioning_error': 10,
    'linear_regression_housing': 11,
    'regression_sensor_calibration': 12,
    'confidence_interval_ab_test': 13,
    'mtbf_estimation': 14,
    'ttest_algorithm_speed': 15,
    'qc_machine_deviation': 16,
    'smart_elevator_queue': 17,
    'bootstrap_resampling': 18
}

topic_dict = {}
with open(topic_file, 'r', encoding='utf-8') as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        try:
            num = int(line.split('.')[0])
            topic_dict[num] = line
        except:
            pass

data = []
index_tuples = []
columns_order = []

files = []
for filename in os.listdir(script_dir):
    if filename.endswith('.json'):
        basename = filename.replace('.json', '')
        idx = filename_to_topic_idx.get(basename, 999)
        files.append((idx, filename))
files.sort(key=lambda x: x[0])

for idx, filename in files:
    basename = filename.replace('.json', '')
    filepath = os.path.join(script_dir, filename)
    topic_name = topic_dict.get(idx, basename)
    
    with open(filepath, 'r', encoding='utf-8') as f:
        try:
            json_data = json.load(f)
        except:
            continue
        
        if isinstance(json_data, list) and len(json_data) > 0:
            item = json_data[0]
            row_voice = {}
            row_visual = {}
            
            for key, content in item.items():
                if key not in columns_order:
                    columns_order.append(key)
                row_voice[key] = content.get("Lời thoại (Voice-over)", "")
                row_visual[key] = content.get("Đoạn video kết quả kèm câu hỏi (Visuals)", "")
            
            data.append(row_voice)
            data.append(row_visual)
            index_tuples.append((topic_name, "Lời thoại (Voice-over)"))
            index_tuples.append((topic_name, "Đoạn video kết quả kèm câu hỏi (Visuals)"))

multi_index = pd.MultiIndex.from_tuples(index_tuples, names=["Chủ đề", "Thành phần"])
df = pd.DataFrame(data, index=multi_index, columns=columns_order)

writer = pd.ExcelWriter(output_file, engine='openpyxl')
df.to_excel(writer, sheet_name='Script')

workbook = writer.book
worksheet = writer.sheets['Script']

worksheet.column_dimensions['A'].width = 80
worksheet.column_dimensions['B'].width = 30

for col_idx in range(3, 3 + len(columns_order)):
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    worksheet.column_dimensions[col_letter].width = 60

for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

for r_idx in range(2, worksheet.max_row + 1):
    worksheet.row_dimensions[r_idx].height = 150

writer.close()
print(f"Excel file created at {output_file} with formatting")
